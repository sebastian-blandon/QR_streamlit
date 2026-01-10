# app.py  —  QR URL/Text + vCard en Streamlit
from __future__ import annotations

import io
import re
import time
import csv
import zipfile
from pathlib import Path
from typing import Optional, Tuple, List, Dict, Any
from urllib.parse import urlparse

import streamlit as st
import streamlit.components.v1 as components
import json
from PIL import Image
import qrcode
from qrcode.constants import ERROR_CORRECT_L, ERROR_CORRECT_M, ERROR_CORRECT_Q, ERROR_CORRECT_H

def set_tab_order(labels_in_order: list[str]) -> None:
    """
    Fija el tabindex de los inputs/textarea de Streamlit según el orden de sus etiquetas.
    Busca por aria-label (Streamlit lo pone igual al texto de la etiqueta).
    """
    js = f"""
    <script>
      const order = {json.dumps(labels_in_order)};
      function applyTabIndex() {{
        let idx = 1;
        order.forEach(lbl => {{
          const el = window.parent.document.querySelector(
            `input[aria-label="${{lbl}}"], textarea[aria-label="${{lbl}}"]`
          );
          if (el) el.tabIndex = idx++;
        }});
      }}
      // Ejecutar tras render; repetir por si hay re-render
      setTimeout(applyTabIndex, 100);
      setTimeout(applyTabIndex, 500);
      setTimeout(applyTabIndex, 1000);
    </script>
    """
    components.html(js, height=0, width=0)

# --- dependencias opcionales (carga perezosa) ---
_HAS_SVG = None
def _get_svg_factory():
    """Devuelve la clase SvgImage o False si no está disponible."""
    global _HAS_SVG
    if _HAS_SVG is None:
        try:
            from qrcode.image.svg import SvgImage
            _HAS_SVG = SvgImage
        except Exception:
            _HAS_SVG = False
    return _HAS_SVG

_HAS_XLSX = None
def _ensure_openpyxl():
    """Devuelve (Workbook, load_workbook) o lanza RuntimeError si falta."""
    global _HAS_XLSX
    if _HAS_XLSX is None:
        try:
            from openpyxl import Workbook, load_workbook
            _HAS_XLSX = (Workbook, load_workbook)
        except Exception:
            _HAS_XLSX = False
    if _HAS_XLSX is False:
        raise RuntimeError(
            "Para usar XLSX instala 'openpyxl':  python -m pip install openpyxl"
        )
    return _HAS_XLSX


# ------------------ Utilidades ------------------

ERROR_MAP = {
    "L (7%)": ERROR_CORRECT_L,
    "M (15%)": ERROR_CORRECT_M,
    "Q (25%)": ERROR_CORRECT_Q,
    "H (30%)": ERROR_CORRECT_H,
}
FORMAT_MAP = ["PNG", "SVG"]


def is_valid_url(url: str) -> bool:
    """Valida mínimamente una URL http(s)."""
    try:
        u = urlparse(url.strip())
        return u.scheme in {"http", "https"} and bool(u.netloc)
    except Exception:
        return False


def is_valid_email(email: str) -> bool:
    """Email opcional: si hay valor, valida patrón básico."""
    if not email.strip():
        return True
    return bool(re.fullmatch(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}", email.strip()))


def is_valid_phone(tel: str) -> bool:
    """Tel opcional: si hay valor, acepta + y espacios/dígitos (>=5)."""
    if not tel.strip():
        return True
    return bool(re.fullmatch(r"\+?[0-9 ]{5,}", tel.strip()))


def slugify(text: str) -> str:
    """Slug simple para nombres de archivo."""
    text = text.strip().lower()
    text = re.sub(r"[^\w\s-]", "", text, flags=re.UNICODE)
    text = re.sub(r"\s+", "_", text)
    return text[:60] or "qr"


def _escape_vcard(value: str) -> str:
    """Escapa caracteres especiales para vCard."""
    value = value.replace("\\", "\\\\").replace("\n", "\\n").replace("\r", "")
    value = value.replace(",", r"\,").replace(";", r"\;")
    return value


def build_payload_url(url: str, title: str, include_title: bool) -> str:
    """Construye el texto para QR en modo URL/Texto."""
    return f"{title.strip()}\n{url.strip()}" if include_title and title.strip() else url.strip()


def generate_qr_pil(
    payload: str,
    version: Optional[int],
    error_label: str,
    box_size: int,
    border: int,
    white_bg: bool = True,
    logo_img: Optional[Image.Image] = None,
    logo_scale_pct: int = 15,
) -> Image.Image:
    """
    Genera un QR como PIL.Image con opción de fondo transparente y logo centrado.
    `logo_img` debe ser PIL.Image (ya cargada); sólo se aplica en PNG/preview.
    """
    version_arg = version if (version and version > 0) else None
    fit_arg = False if version_arg else True
    back_color = "white" if white_bg else "transparent"

    qr = qrcode.QRCode(
        version=version_arg,
        error_correction=ERROR_MAP[error_label],
        box_size=box_size,
        border=border,
    )
    qr.add_data(payload)
    qr.make(fit=fit_arg)

    try:
        img = qr.make_image(fill_color="black", back_color=back_color)
    except Exception:
        img = qr.make_image(fill_color="black", back_color="white")
        if not white_bg:
            img = img.convert("RGBA")
            datas = img.getdata()
            new_data = []
            for p in datas:
                if len(p) == 4:
                    r, g, b, a = p
                    new_data.append((r, g, b, 0) if (r, g, b) == (255, 255, 255) else p)
                else:
                    new_data.append(p)
            img.putdata(new_data)

    if not isinstance(img, Image.Image):
        img = img.get_image()

    if not white_bg and img.mode != "RGBA":
        img = img.convert("RGBA")

    if logo_img:
        try:
            logo = logo_img.convert("RGBA")
            qrw, qrh = img.size
            target = int(min(qrw, qrh) * (logo_scale_pct / 100.0))
            logo.thumbnail((target, target), Image.LANCZOS)
            lw, lh = logo.size
            pos = ((qrw - lw) // 2, (qrh - lh) // 2)
            if img.mode == "RGBA":
                img.alpha_composite(logo, dest=pos)
            else:
                img.paste(logo, pos, mask=logo)
        except Exception:
            pass

    return img


def generate_qr_svg_bytes(
    payload: str,
    version: Optional[int],
    error_label: str,
    box_size: int,
    border: int,
) -> bytes:
    """Genera bytes SVG para el QR (sin logo)."""
    SvgImage = _get_svg_factory()
    if not SvgImage:
        raise RuntimeError("Backend SVG no disponible (instala el extra de svg para 'qrcode').")
    version_arg = version if (version and version > 0) else None
    fit_arg = False if version_arg else True

    qr = qrcode.QRCode(
        version=version_arg,
        error_correction=ERROR_MAP[error_label],
        box_size=box_size,
        border=border,
        image_factory=SvgImage,
    )
    qr.add_data(payload)
    qr.make(fit=fit_arg)
    img = qr.make_image()
    with io.BytesIO() as buf:
        img.save(buf)
        return buf.getvalue()


# ------------------ vCard ------------------

def build_vcard_payload(
    given: str,
    family: str,
    org: str = "",
    title: str = "",
    tel: str = "",
    email: str = "",
    url: str = "",
    street: str = "",
    city: str = "",
    region: str = "",
    postalcode: str = "",
    country: str = "",
    note: str = "",
    version: str = "3.0",
) -> str:
    """
    Construye vCard 3.0/4.0, fijando explícitamente N: Family;Given;Additional;Prefix;Suffix
    Para evitar particiones indeseadas, pon todos los apellidos en 'family' y los nombres en 'given'.
    """
    if not (given.strip() or family.strip()):
        raise ValueError("Debes indicar al menos Nombres o Apellidos para la vCard.")
    if version not in {"3.0", "4.0"}:
        version = "3.0"

    fn = (given.strip() + " " + family.strip()).strip() or (family.strip() or given.strip())

    lines = ["BEGIN:VCARD", f"VERSION:{version}"]
    lines.append(f"N:{_escape_vcard(family)};{_escape_vcard(given)};;;")
    lines.append(f"FN:{_escape_vcard(fn)}")

    if org.strip():
        lines.append(f"ORG:{_escape_vcard(org)}")
    if title.strip():
        lines.append(f"TITLE:{_escape_vcard(title)}")

    if tel.strip():
        tel_esc = _escape_vcard(tel)
        if version == "3.0":
            lines.append(f"TEL;TYPE=CELL:{tel_esc}")
        else:
            lines.append(f"TEL;TYPE=cell:{tel_esc}")

    if email.strip():
        email_esc = _escape_vcard(email)
        if version == "3.0":
            lines.append(f"EMAIL;TYPE=INTERNET:{email_esc}")
        else:
            lines.append(f"EMAIL:{email_esc}")

    if url.strip():
        lines.append(f"URL:{_escape_vcard(url)}")

    if any([street.strip(), city.strip(), region.strip(), postalcode.strip(), country.strip()]):
        adr = f";;{_escape_vcard(street)};{_escape_vcard(city)};{_escape_vcard(region)};{_escape_vcard(postalcode)};{_escape_vcard(country)}"
        if version == "3.0":
            lines.append(f"ADR;TYPE=WORK:{adr}")
        else:
            lines.append(f"ADR;TYPE=work:{adr}")

    if note.strip():
        lines.append(f"NOTE:{_escape_vcard(note)}")

    lines.append("END:VCARD")
    return "\n".join(lines)


# ------------------ XLSX helpers ------------------

URL_HEADERS = ["url", "titulo", "logo"]
VCARD_HEADERS = [
    "nombres", "apellidos", "org", "title", "tel", "email", "url",
    "street", "city", "region", "postalcode", "country", "note", "logo"
]


def read_xlsx_records(xlsx_file: io.BytesIO) -> List[Dict[str, Any]]:
    """Lee la primera hoja de un XLSX subido (BytesIO) y devuelve lista de dicts."""
    _, load_workbook = _ensure_openpyxl()
    wb = load_workbook(xlsx_file, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [(str(h).strip().lower() if h is not None else "") for h in rows[0]]
    recs: List[Dict[str, Any]] = []
    for r in rows[1:]:
        rec = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            val = r[i] if i < len(r) else None
            rec[h] = "" if val is None else str(val).strip()
        if any(v for v in rec.values()):
            recs.append(rec)
    return recs


def write_template_xlsx(mode: str) -> bytes:
    """Devuelve bytes XLSX con plantilla de 'url' o 'vcard' y una fila de ejemplo."""
    Workbook, _ = _ensure_openpyxl()
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla"
    if mode == "url":
        headers = URL_HEADERS
        ws.append(headers)
        ws.append(["https://forms.gle/abc123", "Formulario de inscripción a la conferencia 'Conferencia ABC123'", ""])
    else:
        headers = VCARD_HEADERS
        ws.append(headers)
        ws.append([
            "Sebastián", "Blandón Londoño", "Universidad Tecnológica de Pereira", "Docente Facultad de Ciencias Empresariales",
            "+57 300 000 0000", "s.blandon@utp.edu.co", "https://www.linkedin.com/in/sebastian-blandon/",
            "Carrera 27 #10-02", "Pereira", "Risaralda", "660003", "Colombia",
            "Nota opcional", "C:/Users/Sebas/Pictures/logo.png"
        ])
    for i, h in enumerate(headers, start=1):
        col = ws.column_dimensions[chr(64 + i)]
        col.width = max(12, min(32, len(h) + 2))
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.read()


# ------------------ UI ------------------

st.set_page_config(page_title="Generador de Códigos QR — URL / vCard", layout="wide")

st.title("Generador de Códigos QR — URL / vCard")

with st.sidebar:
    st.header("Parámetros QR")
    ver_choice = st.selectbox("Versión", ["Auto (fit)"] + [str(i) for i in range(1, 41)], index=0)
    version = 0 if ver_choice == "Auto (fit)" else int(ver_choice)
    error_label = st.selectbox("Corrección de error", list(ERROR_MAP.keys()), index=1)
    box_size = st.number_input("Box size", 4, 40, 10, step=1)
    border = st.number_input("Borde", 1, 10, 2, step=1)

    st.header("Exportación")
    fmt = st.radio("Formato", FORMAT_MAP, index=0, horizontal=True)
    white_bg = st.checkbox("Fondo blanco (PNG). Desmarca para transparente", value=True, disabled=(fmt == "SVG"))

    st.header("Logo (PNG/JPG/WebP)")
    logo_file = st.file_uploader("Logo centrado (opcional, sólo PNG)", type=["png", "jpg", "jpeg", "webp"])
    logo_pct = st.slider("Tamaño del logo (% del lado)", min_value=8, max_value=35, value=15)

    logo_img = None
    if logo_file and fmt == "PNG":
        try:
            logo_img = Image.open(logo_file)
            st.caption(f"Logo cargado: {logo_img.size[0]}×{logo_img.size[1]}")
        except Exception as e:
            st.warning(f"No pude abrir el logo: {e}")

tab1, tab2, tab3 = st.tabs(["URL / Texto", "vCard (Contacto)", "Batch XLSX"])

# --------- URL / Texto ---------
with tab1:
    st.subheader("Modo URL / Texto")
    url = st.text_input("URL (https://...)", value="")
    title = st.text_input("Título / Descripción (opcional)", value="")
    include_title = st.checkbox("Incluir título en el contenido del QR (primera línea)", value=True)

    col_a, col_b = st.columns([1,1])
    with col_a:
        if st.button("Previsualizar (URL/Texto)"):
            if not is_valid_url(url):
                st.error("URL inválida. Debe iniciar por http(s)://")
            else:
                payload = build_payload_url(url, title, include_title)
                img = generate_qr_pil(payload, version, error_label, box_size, border, white_bg, logo_img, logo_pct)
                st.image(img, caption="Previsualización", width="stretch")

    with col_b:
        if st.button("Descargar (URL/Texto)"):
            if not is_valid_url(url):
                st.error("URL inválida. Debe iniciar por http(s)://")
            else:
                payload = build_payload_url(url, title, include_title)
                ts = time.strftime("%Y%m%d_%H%M%S")
                base = slugify(title if (include_title and title) else "qr")
                if fmt == "PNG":
                    img = generate_qr_pil(payload, version, error_label, box_size, border, white_bg, logo_img, logo_pct)
                    bio = io.BytesIO()
                    img.save(bio, format="PNG")
                    bio.seek(0)
                    st.download_button("⬇️ Descargar PNG", data=bio.getvalue(),
                                       file_name=f"{base}_{ts}.png", mime="image/png")
                else:
                    try:
                        svg_bytes = generate_qr_svg_bytes(payload, version, error_label, box_size, border)
                        st.download_button("⬇️ Descargar SVG", data=svg_bytes,
                                           file_name=f"{base}_{ts}.svg", mime="image/svg+xml")
                    except Exception as e:
                        st.error(str(e))

    st.divider()
    st.write("🧩 **Plantilla XLSX (modo URL/Texto)**")
    try:
        tpl = write_template_xlsx("url")
        st.download_button("⬇️ Descargar plantilla URL/Texto (XLSX)", data=tpl,
                           file_name="plantilla_url.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        st.warning(str(e))

# --------- vCard ---------
with tab2:
    st.subheader("Modo vCard")
    col1, col2 = st.columns(2)
    with col1:
        given = st.text_input("Nombres", value="")
        org = st.text_input("Organización", value="")
        tel = st.text_input("Teléfono", value="")
        street = st.text_input("Calle", value="")
        region = st.text_input("Región/Depto", value="")
        country = st.text_input("País", value="")
    with col2:
        family = st.text_input("Apellidos", value="")
        title_v = st.text_input("Cargo", value="")
        email = st.text_input("Email", value="")
        city = st.text_input("Ciudad", value="")
        postal = st.text_input("Código Postal", value="")
        url_v = st.text_input("Sitio web", value="")
    note = st.text_input("Nota", value="")
    vver = st.radio("Versión vCard", ["3.0", "4.0"], index=0, horizontal=True)

    col_a, col_b = st.columns([1,1])
    with col_a:
        if st.button("Previsualizar (vCard)"):
            try:
                if email and not is_valid_email(email): raise ValueError("Email inválido.")
                if tel and not is_valid_phone(tel): raise ValueError("Teléfono inválido (usa dígitos y opcional + y espacios).")
                if url_v and not is_valid_url(url_v): raise ValueError("Sitio web inválido en vCard.")
                if not (given or family): raise ValueError("Indica Nombres o Apellidos para la vCard.")
                payload = build_vcard_payload(
                    given=given, family=family, org=org, title=title_v, tel=tel, email=email, url=url_v,
                    street=street, city=city, region=region, postalcode=postal, country=country,
                    note=note, version=vver
                )
                img = generate_qr_pil(payload, version, error_label, box_size, border, white_bg, logo_img, logo_pct)
                st.image(img, caption="Previsualización", width="stretch")
            except Exception as e:
                st.error(str(e))

    with col_b:
        if st.button("Descargar (vCard)"):
            try:
                if email and not is_valid_email(email): raise ValueError("Email inválido.")
                if tel and not is_valid_phone(tel): raise ValueError("Teléfono inválido.")
                if url_v and not is_valid_url(url_v): raise ValueError("Sitio web inválido.")
                if not (given or family): raise ValueError("Indica Nombres o Apellidos para la vCard.")
                payload = build_vcard_payload(
                    given=given, family=family, org=org, title=title_v, tel=tel, email=email, url=url_v,
                    street=street, city=city, region=region, postalcode=postal, country=country,
                    note=note, version=vver
                )
                ts = time.strftime("%Y%m%d_%H%M%S")
                base = slugify((given + " " + family).strip() or (family or given) or "contacto")
                if fmt == "PNG":
                    img = generate_qr_pil(payload, version, error_label, box_size, border, white_bg, logo_img, logo_pct)
                    bio = io.BytesIO()
                    img.save(bio, format="PNG")
                    bio.seek(0)
                    st.download_button("⬇️ Descargar PNG", data=bio.getvalue(),
                                       file_name=f"{base}_{ts}.png", mime="image/png")
                else:
                    svg_bytes = generate_qr_svg_bytes(payload, version, error_label, box_size, border)
                    st.download_button("⬇️ Descargar SVG", data=svg_bytes,
                                       file_name=f"{base}_{ts}.svg", mime="image/svg+xml")
            except Exception as e:
                st.error(str(e))

    st.divider()
    st.write("🧩 **Plantilla XLSX (modo vCard)**")
    try:
        tpl = write_template_xlsx("vcard")
        st.download_button("⬇️ Descargar plantilla vCard (XLSX)", data=tpl,
                           file_name="plantilla_vcard.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        st.warning(str(e))


# Orden deseado izquierda→derecha por fila (ajústalo a tus etiquetas exactas)
set_tab_order([
    "Nombres", "Apellidos",
    "Organización", "Cargo",
    "Teléfono", "Email",
    "Calle", "Ciudad",
    "Región/Depto", "Código Postal",
    "País", "Sitio web",
    "Nota"
])

# --------- Batch XLSX ---------
with tab3:
    st.subheader("Batch desde XLSX (primera hoja)")
    up = st.file_uploader("Sube el XLSX", type=["xlsx"])
    st.caption("Se detecta automáticamente si es URL/Texto o vCard por los encabezados.")
    st.caption("Opcional: sube un logo único que se aplicará a todos (columna 'logo' del XLSX se ignora en modo web).")
    batch_logo = st.file_uploader("Logo para todo el batch (opcional, PNG/JPG/WebP)", type=["png", "jpg", "jpeg", "webp"])
    batch_logo_img = Image.open(batch_logo) if batch_logo else None

    if st.button("Procesar XLSX"):
        try:
            if up is None:
                st.error("Sube primero un archivo XLSX."); st.stop()
            rows = read_xlsx_records(io.BytesIO(up.read()))
            if not rows:
                st.error("XLSX vacío o sin filas útiles."); st.stop()

            headers = set(rows[0].keys())
            is_url = ("url" in headers) and (("titulo" in headers) or ("title" in headers)) and not ({"fn","nombres","given"} & headers)

            ts = time.strftime("%Y%m%d_%H%M%S")
            zbuf = io.BytesIO()
            count_ok, count_fail = 0, 0

            with zipfile.ZipFile(zbuf, "w", compression=zipfile.ZIP_DEFLATED) as z:
                for rec in rows:
                    try:
                        if is_url:
                            urlx = (rec.get("url") or "").strip()
                            titlex = (rec.get("titulo") or rec.get("title") or "").strip()
                            if not is_valid_url(urlx): raise ValueError("URL inválida.")
                            payload = build_payload_url(urlx, titlex, include_title=True)
                            slug = slugify(titlex if titlex else "qr")
                        else:
                            given = (rec.get("given") or rec.get("nombres") or "").strip()
                            family = (rec.get("family") or rec.get("apellidos") or "").strip()
                            fn_csv = (rec.get("fn") or "").strip()
                            if not (given or family):
                                if fn_csv:
                                    family, given = fn_csv, ""
                                else:
                                    raise ValueError("Faltan nombres/apellidos o fn.")
                            org = (rec.get("org") or "").strip()
                            titlev = (rec.get("title") or "").strip()
                            tel = (rec.get("tel") or "").strip()
                            email = (rec.get("email") or "").strip()
                            urlv = (rec.get("url") or "").strip()
                            street = (rec.get("street") or "").strip()
                            city = (rec.get("city") or "").strip()
                            region = (rec.get("region") or "").strip()
                            postal = (rec.get("postalcode") or "").strip()
                            country = (rec.get("country") or "").strip()
                            note = (rec.get("note") or "").strip()

                            if email and not is_valid_email(email): raise ValueError("Email inválido.")
                            if tel and not is_valid_phone(tel): raise ValueError("Teléfono inválido.")
                            if urlv and not is_valid_url(urlv): raise ValueError("URL inválida.")

                            payload = build_vcard_payload(
                                given=given, family=family, org=org, title=titlev, tel=tel, email=email, url=urlv,
                                street=street, city=city, region=region, postalcode=postal, country=country, note=note,
                                version="3.0"
                            )
                            full = (given + " " + family).strip() or family or given or "contacto"
                            slug = slugify(full)

                        if fmt == "PNG":
                            img = generate_qr_pil(
                                payload, 0 if version == 0 else version, error_label, box_size, border,
                                True if white_bg else False,
                                logo_img=(batch_logo_img if batch_logo_img else logo_img),
                                logo_scale_pct=logo_pct
                            )
                            bio = io.BytesIO()
                            img.save(bio, format="PNG")
                            bio.seek(0)
                            z.writestr(f"{slug}_{ts}.png", bio.getvalue())
                        else:
                            svg_bytes = generate_qr_svg_bytes(payload, 0 if version == 0 else version, error_label, box_size, border)
                            z.writestr(f"{slug}_{ts}.svg", svg_bytes)

                        count_ok += 1
                    except Exception:
                        count_fail += 1
                        continue

            zbuf.seek(0)
            st.success(f"Batch finalizado — Éxitos: {count_ok} · Fallos: {count_fail}")
            st.download_button("⬇️ Descargar ZIP", data=zbuf.getvalue(), file_name=f"qr_batch_{ts}.zip", mime="application/zip")
        except Exception as e:
            st.error(str(e))

# Pie de página
st.caption("Nota: el logo sólo se incrusta en PNG. En SVG no se incrusta (limitación del backend).")
